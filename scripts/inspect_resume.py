import docx
from lxml import etree
import re

doc = docx.Document(r'C:\Users\Hanan\OneDrive - University of Saskatchewan\Jobs\Resumes\Hanan_Hussain_Resume_May2026.docx')

def clean_xml(p):
    xml = etree.tostring(p._element, pretty_print=True).decode()
    xml = re.sub(r' xmlns:\w+="[^"]+"', '', xml)
    return xml

for i in [26, 27, 28, 29]:
    print(f'--- Para {i}: {repr(doc.paragraphs[i].text[:60])} ---')
    print(clean_xml(doc.paragraphs[i]))
    print()
